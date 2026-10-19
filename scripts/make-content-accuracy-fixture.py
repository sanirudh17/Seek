"""Known-answer DOCX/PDF parser fixtures, not user documents or rendered artifacts."""
import json
import sys
import zlib
from pathlib import Path
from docx import Document

base = Path(sys.argv[1]).resolve()
base.mkdir(parents=True, exist_ok=True)
body = Document()
p = body.add_paragraph()
p.add_run('quartz')
p.add_run('marker').bold = True
body.add_paragraph('Case CAF\u00c9 and literal a+b are preserved.')
body.add_table(rows=1, cols=1).cell(0, 0).text = 'table cedar phrase'
body.sections[0].header.paragraphs[0].text = 'header copper marker'
body.sections[0].footer.paragraphs[0].text = 'footer silver marker'
body.save(base / 'Structured.docx')
negative = Document()
negative.add_paragraph('No requested term appears in this body.')
negative.save(base / 'quartzmarker-name-only.docx')
(base / 'Broken.docx').write_bytes(b'not a ZIP')
(base / 'Broken.pdf').write_bytes(b'not a PDF')
(base / 'Legacy.doc').write_bytes(b'legacy binary Word not supported')
(base / 'Plain.txt').write_text('quartzmarker in plain text\n', encoding='utf-8')


def pdf(name, pages, title='Owned parser fixture'):
    # Minimal standards-based PDF with real pages/fonts/xref. One text page can
    # contain a compressed stream and split TJ operands, exercising the filter.
    objects = [b'', b'', b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    kids = []
    for commands in pages:
        page_id = len(objects) + 1
        stream_id = page_id + 1
        kids.append(f'{page_id} 0 R')
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 3 0 R >> >> /Contents {stream_id} 0 R >>'.encode())
        data = zlib.compress(commands.encode('ascii'))
        objects.append(f'<< /Length {len(data)} /Filter /FlateDecode >>\nstream\n'.encode() + data + b'\nendstream')
    objects[0] = b'<< /Type /Catalog /Pages 2 0 R >>'
    objects[1] = f'<< /Type /Pages /Kids [{" ".join(kids)}] /Count {len(kids)} >>'.encode()
    objects.append(f'<< /Title ({title}) >>'.encode())
    output = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output += f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n'
    xref = len(output)
    output += f'xref\n0 {len(objects)+1}\n0000000000 65535 f \n'.encode()
    for offset in offsets[1:]:
        output += f'{offset:010d} 00000 n \n'.encode()
    output += f'trailer\n<< /Size {len(objects)+1} /Root 1 0 R /Info {len(objects)} 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    (base / name).write_bytes(output)


pdf('Structured.pdf', ['BT /F1 12 Tf 50 740 Td [(blue) 0 (berry)] TJ 0 -20 Td (quartzmarker PDF body) Tj ET', 'BT /F1 12 Tf 50 740 Td (second page amber phrase) Tj ET'])
pdf('Metadata.pdf', ['BT /F1 12 Tf 50 740 Td (Unrelated visible body) Tj ET'], 'metadataonlytoken')
pdf('NoText.pdf', ['q 0.3 g 50 50 200 200 re f Q'])
cases = []
for query, expected in [
    ('quartzmarker', ['Plain.txt', 'Structured.docx', 'Structured.pdf']),
    ('CAF\u00c9', ['Structured.docx']),
    ('a+b', ['Structured.docx']),
    ('table cedar', ['Structured.docx']),
    ('header copper', ['Structured.docx']),
    ('footer silver', ['Structured.docx']),
    ('blueberry', ['Structured.pdf']),
    ('second page amber', ['Structured.pdf']),
    ('metadataonlytoken', []),
    ('absentneedle', []),
]:
    cases.append({'query': query, 'locations': [str(base)], 'expected': expected})
(base.parent / 'accuracy-manifest.json').write_text(json.dumps(cases), encoding='utf-8')
print(base.parent / 'accuracy-manifest.json')
