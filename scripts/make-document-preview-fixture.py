"""Owned preview fixtures; never opens or changes the user's Seek profile."""
import json
import runpy
import sys
import zipfile
import xml.etree.ElementTree as ET
import io
from PIL import Image, ImageDraw
from pathlib import Path
from openpyxl import Workbook

base = Path(sys.argv[1]).resolve()
base.mkdir(parents=True, exist_ok=True)
sys.argv = ['make-content-accuracy-fixture.py', str(base)]
runpy.run_path(str(Path(__file__).with_name('make-content-accuracy-fixture.py')))
with zipfile.ZipFile(base / 'Structured.docx') as source, zipfile.ZipFile(base / 'TextOnly.docx', 'w', zipfile.ZIP_DEFLATED) as target:
    for entry in source.infolist():
        if entry.filename != 'docProps/thumbnail.jpeg':
            data = source.read(entry)
            if entry.filename == '_rels/.rels':
                relationships = ET.fromstring(data)
                for relationship in list(relationships):
                    if relationship.get('Type', '').endswith('/metadata/thumbnail'):
                        relationships.remove(relationship)
                data = ET.tostring(relationships)
            target.writestr(entry, data)
thumbnail = Image.new('RGB', (400, 520), 'white')
draw = ImageDraw.Draw(thumbnail)
draw.text((30, 35), 'Owned shell-thumbnail fixture', fill='black')
draw.rectangle((30, 90, 370, 160), fill='#3070b0')
image_bytes = io.BytesIO()
thumbnail.save(image_bytes, format='JPEG')
with zipfile.ZipFile(base / 'Structured.docx') as source, zipfile.ZipFile(base / 'Thumbnail.docx', 'w', zipfile.ZIP_DEFLATED) as target:
    for entry in source.infolist():
        target.writestr(entry, image_bytes.getvalue() if entry.filename == 'docProps/thumbnail.jpeg' else source.read(entry))
workbook = Workbook()
sheet = workbook.active
sheet.title = 'First sheet'
sheet.append(['Item', 'Quantity', 'Enabled'])
sheet.append(['Caf\u00e9 quartz', 42, True])
sheet.append(['Second row', 7.5, False])
workbook.create_sheet('Must not show').append(['outside first sheet'])
workbook.save(base / 'Table.xlsx')
large = Workbook()
for n in range(200):
    large.active.append([f'row {n} col {col}' for col in range(15)])
large.save(base / 'Large.xlsx')
Workbook().save(base / 'Empty.xlsx')
(base / 'Broken.xlsx').write_bytes(b'not an Excel archive')
with (base / 'OverLimit.pdf').open('wb') as output:
    output.truncate(32 * 1024 * 1024 + 1)
# pypdf is an optional test-only dependency in an isolated .references folder.
sys.path.insert(0, str(Path('.references/task3-python').resolve()))
from pypdf import PdfReader, PdfWriter
writer = PdfWriter()
for page in PdfReader(base / 'Structured.pdf').pages:
    writer.add_page(page)
writer.encrypt('owned-preview-fixture-only')
with (base / 'Encrypted.pdf').open('wb') as output:
    writer.write(output)
cases = [
    ('Structured.pdf', 0, 'PDF'), ('Structured.pdf', 1, 'PDF'),
    ('Structured.pdf', 2, 'unavailable'), ('NoText.pdf', 0, 'PDF'),
    ('Table.xlsx', 0, 'sheet'), ('Large.xlsx', 0, 'sheet'), ('Empty.xlsx', 0, 'sheet'),
    ('Structured.docx', 0, 'Word'), ('Broken.docx', 0, 'unavailable'),
    ('TextOnly.docx', 0, 'Word text'),
    ('Thumbnail.docx', 0, 'Word thumbnail'),
    ('Broken.pdf', 0, 'unavailable'), ('Broken.xlsx', 0, 'unavailable'),
    ('Encrypted.pdf', 0, 'unavailable'), ('OverLimit.pdf', 0, 'unavailable'),
]
manifest = [{'path': str(base / name), 'page': page, 'expected': expected}
            for name, page, expected in cases]
manifest[4]['contains'] = 'Caf\u00e9 quartz'
manifest.append({'path': str(base / 'Structured.pdf'), 'expected': 'cancelled', 'cancel_after_ms': 1})
(base.parent / 'preview-manifest.json').write_text(json.dumps(manifest), encoding='utf-8')
print(base.parent / 'preview-manifest.json')
